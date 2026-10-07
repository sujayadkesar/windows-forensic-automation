"""Builders for assorted Windows artifacts used by the fixture scenarios."""

from __future__ import annotations

import codecs
import io
import json
import os
import random
import sqlite3
import struct
import uuid
import zipfile
from datetime import datetime, timezone

from .regwriter import filetime
from .shellitems import file_entry, id_list, path_items


# ---------------------------------------------------------------- time utils
def webkit(dt: datetime) -> int:
    return int((dt - datetime(1601, 1, 1, tzinfo=timezone.utc)).total_seconds() * 1_000_000)


def unix(dt: datetime) -> int:
    return int(dt.timestamp())


def prtime(dt: datetime) -> int:
    return int(dt.timestamp() * 1_000_000)


# ---------------------------------------------------------------- recycle bin
def recycle_i_file(original_path: str, size: int, deleted: datetime) -> bytes:
    name = (original_path + "\x00").encode("utf-16-le")
    return struct.pack("<qqqI", 2, size, filetime(deleted), len(original_path) + 1) + name


# ---------------------------------------------------------------- shimcache
def shimcache_win10(entries: list[tuple[str, datetime]]) -> bytes:
    header = bytearray(0x34)
    struct.pack_into("<I", header, 0, 0x34)
    struct.pack_into("<I", header, 0x24, len(entries))
    out = bytes(header)
    for path, ts in entries:
        p = path.encode("utf-16-le")
        data = struct.pack("<H", len(p)) + p + struct.pack("<Q", filetime(ts)) + struct.pack("<I", 0)
        out += struct.pack("<III", 0x73743031, 0, len(data)) + data
    return out


# ---------------------------------------------------------------- userassist
def rot13(s: str) -> str:
    return codecs.encode(s, "rot_13")


def userassist_data(run_count: int, focus_count: int, focus_ms: int, last_run: datetime) -> bytes:
    buf = bytearray(72)
    struct.pack_into("<IIII", buf, 0, 0, run_count, focus_count, focus_ms)
    struct.pack_into("<f", buf, 16, -1.0)
    struct.pack_into("<Q", buf, 60, filetime(last_run))
    return bytes(buf)


# ---------------------------------------------------------------- MRU helpers
def mru_list_ex(n: int) -> bytes:
    return b"".join(struct.pack("<I", i) for i in range(n)) + struct.pack("<I", 0xFFFFFFFF)


def recentdocs_value(name: str, lnk_name: str) -> bytes:
    item = file_entry(lnk_name, is_dir=False)
    return (name + "\x00").encode("utf-16-le") + item + b"\x00\x00"


def opensave_pidl(path: str, times: dict, size: int = 0) -> bytes:
    return id_list(*path_items(path, times, final_is_file=True, size=size))


def lastvisited_pidl(exe: str, folder: str, times: dict) -> bytes:
    return (exe + "\x00").encode("utf-16-le") + id_list(*path_items(folder, times, final_is_file=False))


def bam_value(ts: datetime) -> bytes:
    return struct.pack("<Q", filetime(ts)) + b"\x00" * 8 + struct.pack("<II", 2, 0)


# ---------------------------------------------------------------- chrome
CHROME_SCHEMA = """
CREATE TABLE meta(key LONGVARCHAR NOT NULL UNIQUE PRIMARY KEY, value LONGVARCHAR);
CREATE TABLE urls(id INTEGER PRIMARY KEY AUTOINCREMENT,url LONGVARCHAR,title LONGVARCHAR,visit_count INTEGER DEFAULT 0 NOT NULL,typed_count INTEGER DEFAULT 0 NOT NULL,last_visit_time INTEGER NOT NULL,hidden INTEGER DEFAULT 0 NOT NULL);
CREATE TABLE visits(id INTEGER PRIMARY KEY AUTOINCREMENT,url INTEGER NOT NULL,visit_time INTEGER NOT NULL,from_visit INTEGER,external_referrer_url TEXT,transition INTEGER DEFAULT 0 NOT NULL,segment_id INTEGER,visit_duration INTEGER DEFAULT 0 NOT NULL,incremented_omnibox_typed_score BOOLEAN DEFAULT FALSE NOT NULL,opener_visit INTEGER,originator_cache_guid TEXT,originator_visit_id INTEGER,originator_from_visit INTEGER,originator_opener_visit INTEGER,is_known_to_sync BOOLEAN DEFAULT FALSE NOT NULL,consider_for_ntp_most_visited BOOLEAN DEFAULT FALSE NOT NULL,visited_link_id INTEGER DEFAULT 0 NOT NULL,app_id TEXT);
CREATE TABLE keyword_search_terms (keyword_id INTEGER NOT NULL,url_id INTEGER NOT NULL,term LONGVARCHAR NOT NULL,normalized_term LONGVARCHAR NOT NULL);
CREATE TABLE downloads (id INTEGER PRIMARY KEY,guid VARCHAR NOT NULL,current_path LONGVARCHAR NOT NULL,target_path LONGVARCHAR NOT NULL,start_time INTEGER NOT NULL,received_bytes INTEGER NOT NULL,total_bytes INTEGER NOT NULL,state INTEGER NOT NULL,danger_type INTEGER NOT NULL,interrupt_reason INTEGER NOT NULL,hash BLOB NOT NULL,end_time INTEGER NOT NULL,opened INTEGER NOT NULL,last_access_time INTEGER NOT NULL,transient INTEGER NOT NULL,referrer VARCHAR NOT NULL,site_url VARCHAR NOT NULL,embedder_download_data VARCHAR NOT NULL,tab_url VARCHAR NOT NULL,tab_referrer_url VARCHAR NOT NULL,http_method VARCHAR NOT NULL,by_ext_id VARCHAR NOT NULL,by_ext_name VARCHAR NOT NULL,by_web_app_id VARCHAR NOT NULL,etag VARCHAR NOT NULL,last_modified VARCHAR NOT NULL,mime_type VARCHAR(255) NOT NULL,original_mime_type VARCHAR(255) NOT NULL);
CREATE TABLE downloads_url_chains (id INTEGER NOT NULL,chain_index INTEGER NOT NULL,url LONGVARCHAR NOT NULL, PRIMARY KEY (id, chain_index) );
"""


def chrome_history(path: str, visits: list[dict], downloads: list[dict] | None = None, searches: list[dict] | None = None):
    """visits: dict(url,title,time,transition=0x30000001 typed/link,typed=0)."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        os.remove(path)
    db = sqlite3.connect(path)
    db.executescript(CHROME_SCHEMA)
    db.execute("INSERT INTO meta VALUES ('version','69')")
    url_ids: dict[str, int] = {}
    for v in sorted(visits, key=lambda x: x["time"]):
        url = v["url"]
        if url not in url_ids:
            cur = db.execute(
                "INSERT INTO urls(url,title,visit_count,typed_count,last_visit_time,hidden) VALUES (?,?,?,?,?,0)",
                (url, v.get("title", ""), 0, v.get("typed", 0), webkit(v["time"])),
            )
            url_ids[url] = cur.lastrowid
        uid = url_ids[url]
        db.execute(
            "UPDATE urls SET visit_count=visit_count+1, last_visit_time=?, title=? WHERE id=?",
            (webkit(v["time"]), v.get("title", ""), uid),
        )
        db.execute(
            "INSERT INTO visits(url,visit_time,from_visit,transition,segment_id,visit_duration) VALUES (?,?,0,?,0,?)",
            (uid, webkit(v["time"]), v.get("transition", 0x30000000 | 1), v.get("duration", 15_000_000)),
        )
    for s in searches or []:
        uid = url_ids.get(s["url"])
        if uid:
            db.execute("INSERT INTO keyword_search_terms VALUES (2,?,?,?)", (uid, s["term"], s["term"].lower()))
    for i, d in enumerate(downloads or [], start=1):
        db.execute(
            "INSERT INTO downloads VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                i, str(uuid.uuid4()).upper(), d["target_path"], d["target_path"], webkit(d["start"]),
                d["size"], d["size"], 1, 0, 0, b"", webkit(d["end"]), d.get("opened", 1), webkit(d["end"]), 0,
                d.get("referrer", ""), d.get("site_url", ""), "", d.get("tab_url", ""), d.get("tab_referrer_url", ""),
                "GET", "", "", "", "", "", d.get("mime", "application/octet-stream"), d.get("mime", "application/octet-stream"),
            ),
        )
        db.execute("INSERT INTO downloads_url_chains VALUES (?,0,?)", (i, d["url"]))
    db.commit()
    db.execute("PRAGMA journal_mode=DELETE")
    db.close()


def chrome_webdata(path: str, autofill: list[tuple[str, str, datetime]]):
    if os.path.exists(path):
        os.remove(path)
    db = sqlite3.connect(path)
    db.execute(
        "CREATE TABLE autofill (name VARCHAR, value VARCHAR, value_lower VARCHAR, date_created INTEGER DEFAULT 0, date_last_used INTEGER DEFAULT 0, count INTEGER DEFAULT 1, PRIMARY KEY (name, value))"
    )
    for name, value, ts in autofill:
        db.execute("INSERT INTO autofill VALUES (?,?,?,?,?,1)", (name, value, value.lower(), unix(ts), unix(ts)))
    db.commit()
    db.close()


# ---------------------------------------------------------------- firefox
def firefox_places(path: str, visits: list[dict], downloads: list[dict] | None = None):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        os.remove(path)
    db = sqlite3.connect(path)
    db.executescript(
        """
        CREATE TABLE moz_places (id INTEGER PRIMARY KEY, url LONGVARCHAR, title LONGVARCHAR, rev_host LONGVARCHAR, visit_count INTEGER DEFAULT 0, hidden INTEGER DEFAULT 0 NOT NULL, typed INTEGER DEFAULT 0 NOT NULL, frecency INTEGER DEFAULT -1 NOT NULL, last_visit_date INTEGER , guid TEXT, foreign_count INTEGER DEFAULT 0 NOT NULL, url_hash INTEGER DEFAULT 0 NOT NULL , description TEXT, preview_image_url TEXT, origin_id INTEGER);
        CREATE TABLE moz_historyvisits (id INTEGER PRIMARY KEY, from_visit INTEGER, place_id INTEGER, visit_date INTEGER, visit_type INTEGER, session INTEGER);
        CREATE TABLE moz_anno_attributes (id INTEGER PRIMARY KEY, name VARCHAR(32) UNIQUE NOT NULL);
        CREATE TABLE moz_annos (id INTEGER PRIMARY KEY, place_id INTEGER NOT NULL, anno_attribute_id INTEGER, content LONGVARCHAR, flags INTEGER DEFAULT 0, expiration INTEGER DEFAULT 0, type INTEGER DEFAULT 0, dateAdded INTEGER DEFAULT 0, lastModified INTEGER DEFAULT 0);
        """
    )
    ids: dict[str, int] = {}
    for v in sorted(visits, key=lambda x: x["time"]):
        if v["url"] not in ids:
            host = v["url"].split("/")[2] if "//" in v["url"] else ""
            cur = db.execute(
                "INSERT INTO moz_places(url,title,rev_host,visit_count,typed,last_visit_date,guid) VALUES (?,?,?,0,?,?,?)",
                (v["url"], v.get("title"), host[::-1] + ".", v.get("typed", 0), prtime(v["time"]), uuid.uuid4().hex[:12]),
            )
            ids[v["url"]] = cur.lastrowid
        pid = ids[v["url"]]
        db.execute("UPDATE moz_places SET visit_count=visit_count+1,last_visit_date=? WHERE id=?", (prtime(v["time"]), pid))
        db.execute(
            "INSERT INTO moz_historyvisits(from_visit,place_id,visit_date,visit_type,session) VALUES (0,?,?,?,0)",
            (pid, prtime(v["time"]), v.get("visit_type", 1)),
        )
    if downloads:
        db.execute("INSERT INTO moz_anno_attributes(id,name) VALUES (1,'downloads/destinationFileURI')")
        db.execute("INSERT INTO moz_anno_attributes(id,name) VALUES (2,'downloads/metaData')")
        for d in downloads:
            if d["url"] not in ids:
                cur = db.execute(
                    "INSERT INTO moz_places(url,title,rev_host,visit_count,last_visit_date,guid) VALUES (?,?,?,1,?,?)",
                    (d["url"], os.path.basename(d["target_path"]), "", prtime(d["start"]), uuid.uuid4().hex[:12]),
                )
                ids[d["url"]] = cur.lastrowid
                db.execute(
                    "INSERT INTO moz_historyvisits(from_visit,place_id,visit_date,visit_type,session) VALUES (0,?,?,7,0)",
                    (cur.lastrowid, prtime(d["start"])),
                )
            pid = ids[d["url"]]
            uri = "file:///" + d["target_path"].replace("\\", "/")
            db.execute(
                "INSERT INTO moz_annos(place_id,anno_attribute_id,content,flags,expiration,type,dateAdded,lastModified) VALUES (?,1,?,0,4,3,?,?)",
                (pid, uri, prtime(d["start"]), prtime(d["end"])),
            )
            meta = json.dumps({"state": 1, "endTime": int(d["end"].timestamp() * 1000), "fileSize": d["size"]})
            db.execute(
                "INSERT INTO moz_annos(place_id,anno_attribute_id,content,flags,expiration,type,dateAdded,lastModified) VALUES (?,2,?,0,4,3,?,?)",
                (pid, meta, prtime(d["start"]), prtime(d["end"])),
            )
    db.commit()
    db.close()


# ---------------------------------------------------------------- activities cache
def activities_cache(path: str, activities: list[dict]):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        os.remove(path)
    db = sqlite3.connect(path)
    db.execute(
        """CREATE TABLE [Activity]([Id] GUID PRIMARY KEY NOT NULL, [AppId] TEXT NOT NULL, [PackageIdHash] TEXT,
        [AppActivityId] TEXT, [ActivityType] INT NOT NULL, [ActivityStatus] INT NOT NULL, [ParentActivityId] GUID,
        [Tag] TEXT, [Group] TEXT, [MatchId] TEXT, [LastModifiedTime] DATETIME NOT NULL, [ExpirationTime] DATETIME,
        [Payload] BLOB, [Priority] INT, [IsLocalOnly] INT, [PlatformDeviceId] TEXT, [CreatedInCloud] DATETIME,
        [StartTime] DATETIME, [EndTime] DATETIME, [LastModifiedOnClient] DATETIME, [GroupAppActivityId] TEXT,
        [ClipboardPayload] BLOB, [EnterpriseId] TEXT, [OriginalPayload] BLOB, [OriginalLastModifiedOnClient] DATETIME,
        [ETag] INT NOT NULL)"""
    )
    for a in activities:
        app_id = json.dumps([{"application": a["app"], "platform": "x_exe_path"}, {"application": a["app_name"], "platform": "windows_win32"}])
        payload = json.dumps(
            {"displayText": a["display"], "activationUri": a.get("uri", ""), "appDisplayName": a.get("app_display", ""),
             "description": a.get("description", ""), "contentUri": a.get("uri", "")}
        ).encode()
        start = unix(a["start"])
        end = unix(a.get("end", a["start"]))
        db.execute(
            "INSERT INTO Activity VALUES (?,?,?,?,?,?,NULL,NULL,NULL,NULL,?,?,?,3,0,'',0,?,?,?,NULL,?,'',NULL,0,1)",
            (uuid.uuid4().bytes, app_id, "", a.get("uri", ""), a.get("type", 5), 1, start, start + 30 * 86400,
             payload, start, end, start, a.get("clipboard")),
        )
    db.commit()
    db.close()


# ---------------------------------------------------------------- documents
def make_docx(path: str, title: str, paragraphs: list[str], author: str = "") -> None:
    import docx  # python-docx

    d = docx.Document()
    d.add_heading(title, 0)
    for p in paragraphs:
        d.add_paragraph(p)
    d.core_properties.author = author
    d.core_properties.last_modified_by = author
    d.core_properties.title = title
    d.save(path)


def make_xlsx(path: str, rows: list[list], author: str = "") -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.properties.creator = author
    wb.save(path)


def make_pdf(path: str, lines: list[str]) -> None:
    """Very small single page PDF with uncompressed text (searchable)."""
    text_ops = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(
        "(" + ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '" for ln in lines
    ) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(text_ops)} >>\nstream\n{text_ops}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        "<< /Title (Merger Term Sheet) /Author (Legal Dept) >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info 6 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    with open(path, "wb") as fh:
        fh.write(out.getvalue())


def make_zip(path: str, files: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


# ---------------------------------------------------------------- memory-ish blobs
def noise_blob(size: int, embedded: list[tuple[int, bytes]], seed: int = 7) -> bytes:
    rnd = random.Random(seed)
    buf = bytearray(rnd.getrandbits(8) for _ in range(min(size, 1 << 16)))
    while len(buf) < size:
        buf += buf[: min(len(buf), size - len(buf))]
    for off, data in embedded:
        buf[off:off + len(data)] = data
    return bytes(buf[:size])


def utf16(s: str) -> bytes:
    return s.encode("utf-16-le")


def filetime_bytes(dt: datetime) -> bytes:
    return struct.pack("<Q", filetime(dt))
