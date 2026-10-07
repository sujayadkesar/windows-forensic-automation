"""Low level python-docx helpers (shading, borders, fields, table formatting)."""

from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

NAVY = "1F3864"
NAVY_RGB = RGBColor(0x1F, 0x38, 0x64)
GREY_RGB = RGBColor(0x59, 0x59, 0x59)

STATUS_COLORS = {  # conclusion status -> (fill, text); deliberately neutral - this is an evidential report, not a risk rating
    "Yes": ("D6E4F5", "1F3864"), "Indicated": ("E9F0F9", "2E5C8A"), "No evidence found": ("F2F2F2", "404040"),
    "Not applicable": ("FAFAFA", "7F7F7F"), "Inconclusive": ("FFF8E5", "7F6000"),
}
STATUS_ACCENT = {"Yes": "1F3864", "Indicated": "2E75B6", "No evidence found": "7F7F7F", "Not applicable": "A6A6A6",
                 "Inconclusive": "BF9000"}
STATUS_ACCENT = {"Yes": "1F3864", "Indicated": "2E75B6", "No evidence found": "7F7F7F", "Not applicable": "A6A6A6",
                 "Inconclusive": "BF9000"}
SEV_COLORS = {"critical": ("7B0000", "FFFFFF"), "high": ("C00000", "FFFFFF"), "medium": ("ED7D31", "FFFFFF"),
              "low": ("FFD966", "3F3F00"), "info": ("BDD7EE", "1F3864")}
COV_COLORS = {"found": ("C6EFCE", "006100"), "not_found": ("F2F2F2", "404040"), "absent": ("EDEDED", "7F7F7F"),
              "error": ("FFC7CE", "9C0006"), "skipped": ("FFF2CC", "7F6000"), "partial": ("FFEB9C", "9C5700")}
COV_LABEL = {"found": "Found", "not_found": "Checked - nothing found", "absent": "Not present on system", "error": "Error",
             "skipped": "Skipped", "partial": "Partial"}


def shade(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_cell_text(cell, text, bold=False, size=8.5, color: str | None = None, align=None, italic=False, font=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if font:
        run.font.name = font
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def table_borders(table, color="BFBFBF", size=4, inside=True):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ())
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_col_widths(table, widths_twips: list[int]):
    """Fixed layout + explicit grid widths (python-docx needs both cell and grid widths)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_twips):
            gc.set(qn("w:w"), str(widths_twips[i]))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_twips):
                tcPr = cell._tc.get_or_add_tcPr()
                w = tcPr.find(qn("w:tcW"))
                if w is None:
                    w = OxmlElement("w:tcW")
                    tcPr.append(w)
                w.set(qn("w:w"), str(widths_twips[i]))
                w.set(qn("w:type"), "dxa")


def add_field(paragraph, instr: str, placeholder: str = ""):
    """Insert a Word field (PAGE, NUMPAGES, TOC ...)."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    b.set(qn("w:dirty"), "true")
    run._r.append(b)
    run2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run2._r.append(it)
    run3 = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(sep)
    run4 = paragraph.add_run(placeholder)
    run5 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run5._r.append(e)
    return run4


def update_fields_on_open(document):
    settings = document.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def paragraph_shading(paragraph, fill_hex: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def paragraph_border_left(paragraph, color="1F3864", size=24):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pbdr.append(left)
    pPr.append(pbdr)


def keep_with_next(paragraph, on=True):
    paragraph.paragraph_format.keep_with_next = on


def center_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def bookmark(paragraph, name: str, bid: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def cant_split(row) -> None:
    """Keep a table row on one page."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)
