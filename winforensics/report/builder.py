"""Professional forensic report (Word) + Excel workbook + figures.

``build_report(case)`` renders every finding figure as an annotated screenshot,
writes the Excel workbook of all parsed artifacts, composes the Word report and -
when Microsoft Word is installed - lets Word update the table of contents and
export a PDF copy.
"""

from __future__ import annotations

import json
import os
import platform
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from .. import __app_title__, __version__
from ..core.timeutil import UTC, fmt, fmt_dual, get_tz
from . import docx_util as U

CONTENT_W_CM = 17.0
TW = 567  # twips per cm


def _cm(x):
    return int(x * TW)


class ReportWriter:
    def __init__(self, case, figure_paths: dict, out_path: str, progress=None):
        from ..profiles import get_profile

        self.case = case
        self.db = case.db
        self.info = case.info
        self.profile = get_profile(self.info.profile)
        self.inputs = self.db.meta("inputs_normalized") or {}
        self.evidence = self.db.evidence()
        self.findings = [f for f in self.db.findings(include_only=True)]
        self.answers = self.db.answers()
        self.types = self.db.artifact_types()
        self.tz = get_tz(self.info.display_timezone)
        self.fig_paths = figure_paths
        self.out_path = out_path
        self.progress = progress or (lambda f, s: None)
        self.fig_no = 0
        self.fnum = {}
        self.doc = Document()

    # ================================================================== helpers
    def ev_label(self, eid):
        for e in self.evidence:
            if e["id"] == eid:
                host = (e.get("os") or {}).get("hostname")
                return e["label"] + (f" ({host})" if host and host not in e["label"] else "")
        return "Case-wide"

    def t(self, ts):
        return fmt_dual(ts, self.tz) if ts else "-"

    def heading(self, text, level=1):
        h = self.doc.add_heading(text, level=level)
        U.keep_with_next(h)
        if level == 1 and getattr(self, "_break_next", False):
            # 'page break before' instead of a break paragraph: never produces an empty page after a full one
            h.paragraph_format.page_break_before = True
            self._break_next = False
        return h

    def new_page(self):
        self._break_next = True

    def para(self, text="", bold=False, italic=False, size=None, color=None, align=None, space_after=6, style=None):
        p = self.doc.add_paragraph(style=style)
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def rich(self, parts, style=None, space_after=6):
        """parts: [(text, {bold, italic, color})]"""
        p = self.doc.add_paragraph(style=style)
        for text, f in parts:
            r = p.add_run(text)
            r.bold = f.get("bold", False)
            r.italic = f.get("italic", False)
            if f.get("color"):
                r.font.color.rgb = RGBColor.from_string(f["color"])
            if f.get("size"):
                r.font.size = Pt(f["size"])
            if f.get("font"):
                r.font.name = f["font"]
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def bullet(self, text, bold_prefix=""):
        p = self.doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            p.add_run(bold_prefix).bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(2)
        return p

    def table(self, headers, rows, widths_cm, font_size=8.5, header_fill=U.NAVY, zebra=True, status_col=None,
              status_map=None, mono_cols=()):
        t = self.doc.add_table(rows=1, cols=len(headers))
        U.table_borders(t, "BFBFBF", 4)
        hdr = t.rows[0]
        U.repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            U.shade(c, header_fill)
            U.set_cell_text(c, h, bold=True, size=font_size, color="FFFFFF")
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for c in hdr.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
        for r_i, row in enumerate(rows):
            new_row = t.add_row()
            U.cant_split(new_row)
            cells = new_row.cells
            for i, v in enumerate(row):
                txt = "" if v is None else str(v)
                if status_col is not None and i == status_col and status_map:
                    fill, fg = status_map.get(txt, ("FFFFFF", "000000"))
                    label = U.COV_LABEL.get(txt, txt) if status_map is U.COV_COLORS else txt
                    U.shade(cells[i], fill)
                    U.set_cell_text(cells[i], label, bold=True, size=font_size - 0.5, color=fg)
                else:
                    U.set_cell_text(cells[i], txt, size=font_size, font="Consolas" if i in mono_cols else None)
                    if zebra and r_i % 2:
                        U.shade(cells[i], "F2F5FA")
        U.set_col_widths(t, [_cm(w) for w in widths_cm])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, rows, widths=(4.2, 12.8), label_fill="EEF2F8"):
        t = self.doc.add_table(rows=0, cols=2)
        U.table_borders(t, "D0D7E5", 4)
        for k, v in rows:
            cells = t.add_row().cells
            U.shade(cells[0], label_fill)
            U.set_cell_text(cells[0], k, bold=True, size=9, color="1F3864")
            U.set_cell_text(cells[1], v, size=9)
        U.set_col_widths(t, [_cm(w) for w in widths])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, path, title, caption=""):
        if not path or not os.path.exists(path):
            return
        from PySide6.QtGui import QImage

        from .figures import S

        img = QImage(path)
        logical = img.width() / S if not img.isNull() else 1100
        # ~150 logical px per inch keeps UI text at a legible print size; never wider than the text column
        width = Cm(min(CONTENT_W_CM, logical / 150 * 2.54))
        self.doc.add_picture(path, width=width)
        pic_par = self.doc.paragraphs[-1]
        pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_par.paragraph_format.left_indent = Cm(0)
        pic_par.paragraph_format.right_indent = Cm(0)
        pic_par.paragraph_format.first_line_indent = Cm(0)
        pic_par.paragraph_format.space_before = Pt(4)
        pic_par.paragraph_format.keep_with_next = True
        pic_par.paragraph_format.space_after = Pt(2)
        self.fig_no += 1
        cap = self.doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.fig_no}: ")
        r.bold = True
        cap.add_run(title + (f". {caption}" if caption else ""))
        cap.paragraph_format.space_after = Pt(10)

    # ================================================================== document setup
    def setup(self):
        d = self.doc
        st = d.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1
        for name, size, color, before in (("Heading 1", 16, U.NAVY_RGB, 18), ("Heading 2", 13, U.NAVY_RGB, 14),
                                          ("Heading 3", 11.5, RGBColor(0x2E, 0x54, 0x96), 12)):
            s = st[name]
            s.font.name = "Calibri"
            rf = s.element.rPr.rFonts
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(att)) is not None:
                    del rf.attrib[qn(att)]
            rf.set(qn("w:ascii"), "Calibri")
            rf.set(qn("w:hAnsi"), "Calibri")
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = color
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(6)
        cap = st["Caption"]
        cap.font.size = Pt(8.5)
        cap.font.italic = False
        cap.font.color.rgb = U.GREY_RGB
        sec = d.sections[0]
        paper = (self.info.report or {}).get("paper", "A4")
        if paper == "Letter":
            sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
        else:
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(1.8)
        sec.header_distance = Cm(0.9)
        sec.footer_distance = Cm(0.8)
        sec.different_first_page_header_footer = True
        # running header / footer
        hp = sec.header.paragraphs[0]
        self._clear_style_tabs(hp)
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run(f"{self.info.case_number or self.case.title}  |  {self.profile.title}")
        r.font.size = Pt(8)
        r.font.color.rgb = U.GREY_RGB
        r = hp.add_run(f"\t{self.info.classification or 'CONFIDENTIAL'}")
        r.font.size = Pt(8)
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        fp = sec.footer.paragraphs[0]
        self._clear_style_tabs(fp)
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
        r = fp.add_run(f"{__app_title__} {__version__}  -  generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC")
        r.font.size = Pt(8)
        r.font.color.rgb = U.GREY_RGB
        r = fp.add_run("\tPage ")
        r.font.size = Pt(8)
        U.add_field(fp, "PAGE", "1").font.size = Pt(8)
        r = fp.add_run(" of ")
        r.font.size = Pt(8)
        U.add_field(fp, "NUMPAGES", "1").font.size = Pt(8)

    @staticmethod
    def _clear_style_tabs(paragraph):
        # the built-in Header / Footer styles define center (3.25in) and right (6.5in) tab stops
        from docx.shared import Inches

        for pos in (Inches(3.25), Inches(6.5)):
            paragraph.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)

    # ================================================================== sections
    def cover(self, logo_path):
        d = self.doc
        t = d.add_table(rows=1, cols=2)
        U.no_borders(t)
        c0, c1 = t.rows[0].cells
        c0.paragraphs[0].add_run().add_picture(logo_path, width=Cm(1.6))
        r = c0.paragraphs[0].add_run("   WINDOWS FORENSIC AUTOMATION")
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = U.NAVY_RGB
        U.shade(c1, "C00000")
        U.set_cell_text(c1, self.info.classification or "CONFIDENTIAL", bold=True, size=10, color="FFFFFF", align="center")
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        U.set_col_widths(t, [_cm(12.5), _cm(4.5)])
        for _ in range(4):
            d.add_paragraph()
        band = d.add_table(rows=1, cols=1)
        U.no_borders(band)
        cell = band.rows[0].cells[0]
        U.shade(cell, "1F3864")
        U.cell_margins(cell, 360, 360, 360, 360)
        p = cell.paragraphs[0]
        r = p.add_run(self.profile.title)
        r.bold = True
        r.font.size = Pt(26)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p2 = cell.add_paragraph()
        r = p2.add_run(f"Investigation profile: {self.profile.name}")
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0x9D, 0xC3, 0xE6)
        if self.info.case_name:
            p3 = cell.add_paragraph()
            r = p3.add_run(self.info.case_name)
            r.font.size = Pt(15)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        U.set_col_widths(band, [_cm(CONTENT_W_CM)])
        d.add_paragraph()
        d.add_paragraph()
        rows = [("Case number", self.info.case_number or "-"), ("Case name", self.info.case_name or "-"),
                ("Client / requester", self.info.client or "-"), ("Examiner", self.info.examiner or "-"),
                ("Organization", self.info.organization or "-"), ("Report date", datetime.now(UTC).strftime("%d %B %Y")),
                ("Evidence items", ", ".join(self.ev_label(e["id"]) for e in self.evidence) or "-"),
                ("Time zone", f"All times UTC; local times shown in {self.info.display_timezone}")]
        self.kv_table(rows, (5.0, 12.0))
        for _ in range(3):
            d.add_paragraph()
        p = self.para("This report was produced with the assistance of automated analysis. All findings were generated from the "
                      "examined evidence and must be reviewed and validated by the examiner before the report is relied upon.",
                      italic=True, size=8.5, color="595959")
        d.add_page_break()

    def document_control(self):
        self.heading("Document control", 1)
        self.table(["Version", "Date", "Author", "Status", "Comment"],
                   [["1.0", datetime.now(UTC).strftime("%Y-%m-%d"), self.info.examiner or "-", "Draft - examiner review",
                     f"Generated by {__app_title__} {__version__}"]], [1.6, 2.4, 4.0, 3.8, 5.2])
        self.heading("Contents", 1)
        p = self.doc.add_paragraph()
        U.add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose 'Update Field' to build the table of contents.")
        self.new_page()

    def executive_summary(self):
        self.heading("1. Executive summary", 1)
        evs = ", ".join(self.ev_label(e["id"]) for e in self.evidence)
        bg = (self.inputs.get("background") or self.info.description or "").strip()
        text = (f"{self.info.organization or 'The examiner'} examined {len(self.evidence)} evidence item(s) - {evs} - using the "
                f"'{self.profile.name}' investigation profile of {__app_title__}. ")
        if bg:
            text += f"Background: {bg} "
        win = self.inputs.get("time_window") or {}
        if win.get("start") or win.get("end"):
            text += (f"The period of interest was {fmt(win.get('start'), UTC)} to {fmt(win.get('end'), UTC)}"
                     + (" (derived from the DLP alerts)" if win.get("derived") else "") + ". ")
        self.para(text)
        questions = self.profile.questions
        positives = [(i, q) for i, q in enumerate(questions, 1) if self.answers.get(q["id"], {}).get("status") in ("Yes", "Indicated")]
        negatives = [(i, q) for i, q in enumerate(questions, 1) if self.answers.get(q["id"], {}).get("status") == "No evidence found"]
        self.heading("Evidence examined", 2)
        erows = []
        for e in self.evidence:
            osd = e.get("os") or {}
            hashes = e.get("hashes") or {}
            ver = ("verified" if hashes.get("verified") else "MISMATCH" if hashes.get("verified") is False else
                   "computed" if hashes.get("md5") else "container checks only")
            erows.append([e["label"], self.profile.role(e["role"]).get("label", e["role"]), osd.get("hostname") or "-",
                          (osd.get("product_name") or "-") + (f" ({osd.get('build')})" if osd.get("build") else ""),
                          f"{((e.get('info') or {}).get('media_size') or e.get('size') or 0) / 1e9:.1f} GB", ver])
        self.table(["Evidence", "Role", "Host", "Operating system", "Size", "Integrity"], erows, [3.0, 3.6, 2.6, 3.8, 1.6, 2.4],
                   font_size=8)
        if positives:
            self.heading("Key conclusions", 2)
            for i, q in positives:
                a = self.answers[q["id"]]
                p = self.doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(f"Q{i} - {a['status']}: ")
                r.bold = True
                p.add_run(q["text"])
                for line in [x.strip() for x in (a.get("summary") or "").split("\n") if x.strip()][:6]:
                    sp = self.doc.add_paragraph(style="List Bullet 2")
                    sp.paragraph_format.space_after = Pt(1)
                    sp.add_run(line[:400]).font.size = Pt(9)
        if negatives:
            self.heading("No evidence found for", 2)
            for i, q in negatives:
                self.bullet(q["text"], bold_prefix=f"Q{i}: ")
        self.heading("Investigative questions and conclusions", 2)
        rows = []
        for i, q in enumerate(questions, 1):
            a = self.answers.get(q["id"], {})
            refs = sorted({self.fnum[x] for x in a.get("finding_ids", []) if x in self.fnum})
            summary = " ".join(x.strip() for x in (a.get("summary") or "Not assessed").split("\n"))[:700] + \
                (f"  (See {', '.join(refs[:8])})" if refs else "")
            rows.append([f"Q{i}", q["text"], a.get("status", "Inconclusive"), summary])
        self.table(["#", "Question", "Conclusion", "Basis"], rows, [1.0, 5.3, 2.4, 8.3], status_col=2, status_map=U.STATUS_COLORS)
        self.para("Conclusion key: 'Yes' - established by the artifacts; 'Indicated' - supported by artifacts but not conclusive on "
                  "its own; 'No evidence found' - the relevant artifacts were examined and nothing was found (absence of evidence is "
                  "not evidence of absence - see the coverage matrix in section 8); 'Not applicable' - the question could not apply "
                  "to the evidence or inputs supplied.", italic=True, size=8.5, color="595959")
        self.new_page()

    def background(self):
        self.heading("2. Case background and examiner inputs", 1)
        bg = (self.inputs.get("background") or self.info.description or "").strip()
        self.para(bg or "No narrative background was supplied.")
        sub = iter(range(1, 10))
        self.heading(f"2.{next(sub)} Scope", 2)
        self.para(f"Profile: {self.profile.name}. {self.profile.description.strip()}")
        dlp = self.inputs.get("dlp_events") or []
        if dlp:
            self.heading(f"2.{next(sub)} DLP alerts supplied", 2)
            rows = [[i + 1, (e.get("time") or "")[:19], e.get("activity"), e.get("user"), e.get("device"), e.get("file_name"),
                     e.get("usb_serial") or e.get("destination") or ""] for i, e in enumerate(dlp[:200])]
            self.table(["#", "Time (UTC)", "Activity", "User", "Device", "File", "USB serial / destination"], rows,
                       [0.7, 2.8, 3.3, 1.8, 2.4, 3.2, 2.8], font_size=7.5)
        tg = self.inputs.get("targets") or []
        if tg:
            self.heading(f"2.{next(sub)} Files of interest", 2)
            rows = [[i + 1, t.get("name"), t.get("size") or "", (t.get("sha256") or t.get("md5") or "")[:64], ", ".join(t.get("sources") or [])]
                    for i, t in enumerate(tg[:300])]
            self.table(["#", "File name", "Size", "Hash (SHA-256 / MD5)", "Source of information"], rows, [0.7, 4.3, 1.6, 7.0, 3.4],
                       font_size=7.5, mono_cols=(3,))
        other = []
        for k, label in (("usb_serials", "USB serial numbers"), ("users", "Users of interest"), ("devices", "Devices named in alerts"),
                         ("domains", "Destinations of concern")):
            if self.inputs.get(k):
                other.append((label, ", ".join(self.inputs[k])))
        win = self.inputs.get("time_window") or {}
        if win:
            other.append(("Period of interest (UTC)", f"{fmt(win.get('start'), UTC)} - {fmt(win.get('end'), UTC)}" +
                          (" (derived from DLP alerts +/- 24 h)" if win.get("derived") else "")))
        kws = self.inputs.get("keywords") or []
        if kws:
            other.append(("Search terms", f"{len(kws)} terms: " + ", ".join(k['term'] for k in kws[:40]) + (" ..." if len(kws) > 40 else "")))
        if other:
            self.heading(f"2.{next(sub)} Other inputs", 2)
            self.kv_table(other)

    def evidence_section(self):
        self.heading("3. Evidence examined", 1)
        self.para("Each evidence item was opened read-only. Container metadata and, where requested, a full verification hash were "
                  "recorded. The operating system details below were read from the image itself.")
        for e in self.evidence:
            info = e.get("info") or {}
            osd = e.get("os") or {}
            ewf = info.get("ewf") or {}
            hashes = e.get("hashes") or {}
            vols = info.get("volumes") or []
            rows = [("Label / role", f"{e['label']}  -  {self.profile.role(e['role']).get('label', e['role'])}"),
                    ("Image file(s)", e["path"] + (f"  ({len(ewf.get('segments', []))} segments)" if ewf.get("segments") else "")),
                    ("Format", f"{e.get('format')}"), ("Media size", f"{(info.get('media_size') or 0) / 1e9:.2f} GB"),
                    ]
            if ewf.get("headers"):
                h = ewf["headers"]
                rows.append(("Acquisition metadata", "; ".join(f"{k}: {v}" for k, v in h.items() if k in (
                    "case_number", "evidence_number", "examiner_name", "acquiry_date", "acquiry_software_version", "description"))))
            if ewf.get("stored_hashes"):
                rows.append(("Acquisition hash", ", ".join(f"{k.upper()} {v}" for k, v in ewf["stored_hashes"].items())))
            if hashes.get("md5"):
                ver = hashes.get("verified")
                rows.append(("Verification", f"MD5 {hashes.get('md5')}  SHA1 {hashes.get('sha1')}  -  " +
                             ("VERIFIED (matches acquisition hash)" if ver else "MISMATCH" if ver is False else "computed")))
            else:
                rows.append(("Verification", "Not re-verified during this examination"))
            if osd.get("family") == "windows":
                rows += [("Host name", osd.get("hostname") or osd.get("computer_name") or "-"),
                         ("Operating system", f"{osd.get('product_name') or ''} {osd.get('display_version') or ''} (build {osd.get('build') or '?'})"),
                         ("Installed", fmt(osd.get("install_date"), UTC) or "-"),
                         ("Time zone", f"{osd.get('timezone_name') or '-'} ({osd.get('timezone_iana') or ''})"),
                         ("Last shutdown", fmt(osd.get("last_shutdown"), UTC) or "-"),
                         ("User profiles", ", ".join(u.get("name", "") for u in info.get("users", [])) or "-")]
            vtxt = []
            for v in vols:
                if v.get("fs"):
                    vtxt.append(f"{v.get('name')}: {str(v.get('fs')).upper()} '{v.get('label') or ''}' serial {v.get('serial') or '-'} "
                                f"({(v.get('size') or 0) / 1e9:.2f} GB)")
            rows.append(("Volumes", "\n".join(vtxt) or "-"))
            self.heading(f"3.{self.evidence.index(e) + 1} {self.ev_label(e['id'])}", 2)
            self.kv_table(rows)

    def methodology(self):
        from ..modules.base import MODULES, discover

        discover()
        self.heading("4. Methodology", 1)
        self.para(f"The evidence was processed with {__app_title__} {__version__}, an automated, profile-driven Windows forensic "
                  "platform. Each image was opened read-only; no data was written to the evidence. The following steps were "
                  "performed for every Windows system image:")
        ran = sorted({r["module"] for r in self.db.query("SELECT DISTINCT module FROM coverage")})
        for m in ran:
            cls = MODULES.get(m)
            if cls:
                self.bullet(cls.description, bold_prefix=f"{cls.title}: ")
        self.heading("4.1 Search methodology", 2)
        self.para("Physical keyword search reads every byte of every volume and of unpartitioned disk space and matches every term in "
                  "ASCII/UTF-8 and UTF-16LE, case-insensitively. Each hit is attributed using the file system's own allocation "
                  "structures to an allocated file, file slack, a deleted file's former clusters, unallocated space, an MFT record, the "
                  "$LogFile, the USN journal, a directory index, pagefile.sys / hiberfil.sys / swapfile.sys or a shadow copy store. "
                  "Because the content of compressed documents (DOCX, XLSX, PDF) is not visible to a physical search, the text of every "
                  "document - including documents recovered from deleted entries and carved from unallocated space - was also "
                  "indexed and searched. Files of interest were located by MD5/SHA-1/SHA-256, by name and size, by text similarity and "
                  "by document metadata, in allocated and deleted files, the Recycle Bin, ZIP archives, e-mail attachments, carved "
                  "files and Volume Shadow Copies.")
        self.heading("4.2 Time", 2)
        self.para(f"All timestamps are stored and reported in UTC; where useful the local time in {self.info.display_timezone} is shown "
                  "in brackets. Artifacts that record local time (setupapi logs, FAT directory entries, NetworkList dates) were "
                  "converted using the time zone configured on the source system, or the assumption stated in the limitations.")

    def findings_section(self):
        self.heading("5. Findings", 1)
        qs = self.profile.questions
        placed = set()
        order = []
        for i, q in enumerate(qs, 1):
            fs = [f for f in self.findings if q["id"] in (f.get("questions") or []) and f["id"] not in placed]
            for f in fs:
                placed.add(f["id"])
            order.append((f"5.{i} Q{i}: {q['text']}", q, fs))
        rest = [f for f in self.findings if f["id"] not in placed]
        if rest:
            order.append((f"5.{len(qs) + 1} Other observations", None, rest))
        total = sum(len(x[2]) for x in order) or 1
        done = 0
        for title, q, fs in order:
            self.heading(title, 2)
            if q is not None:
                a = self.answers.get(q["id"], {})
                st = a.get("status", "Inconclusive")
                summ = " ".join(x.strip() for x in (a.get("summary") or "").split("\n")).strip() or \
                    "No analyzer reached a conclusion for this question with the evidence available; see the findings and the coverage matrix."
                p = self.rich([("Conclusion: ", {"bold": True}), (st, {"bold": True, "color": U.STATUS_ACCENT.get(st, "1F3864")}),
                               (" - " + summ[:900], {})])
                U.paragraph_border_left(p, U.STATUS_ACCENT.get(st, "1F3864"), 18)
                U.paragraph_shading(p, "F7F9FC")
            if not fs:
                self.para("No specific findings were produced for this question; see the conclusion above and the coverage matrix.",
                          italic=True, color="595959")
            for f in fs:
                self.finding(f)
                done += 1
                self.progress(0.55 + 0.3 * done / total, f"Writing finding {done}/{total}")

    def finding(self, f):
        num = self.fnum[f["id"]]
        self.heading(f"{num}  {f['title']}", 3)
        t = self.doc.add_table(rows=1, cols=3)
        U.table_borders(t, "D9D9D9", 4)
        cells = t.rows[0].cells
        nrefs = len([r for r in (f.get("refs") or []) if r.get("kind") == "artifact"])
        U.set_cell_text(cells[0], "System: " + self.ev_label(f.get("evidence_id")), size=8, align="center")
        U.set_cell_text(cells[1], "Time: " + (fmt(f["ts"], UTC) if f.get("ts") else "-"), size=8, align="center")
        U.set_cell_text(cells[2], f"Supporting artifacts: {nrefs}" + (f"   |   {f.get('category')}" if f.get("category") else ""),
                        size=8, align="center")
        for c in cells:
            U.shade(c, "F2F5FA")
        U.set_col_widths(t, [_cm(6.2), _cm(5.0), _cm(5.8)])
        self.para("", space_after=2)
        self.para(f["description"])
        for i, fig in enumerate(f.get("figures") or []):
            path = self.fig_paths.get((f["id"], i))
            self.figure(path, fig.get("title", ""), fig.get("caption", ""))
        refs = [r for r in (f.get("refs") or []) if r.get("kind") == "artifact"][:8]
        if refs:
            rows = []
            for r in refs:
                a = self.db.query("SELECT type, ts, summary, source FROM artifacts WHERE id=?", (r["id"],))
                if a:
                    a = a[0]
                    rows.append([self.types.get(a["type"], {}).get("title", a["type"]), fmt(a["ts"], UTC, with_zone=False) if a["ts"] else "",
                                 (a["summary"] or "")[:160], (a["source"] or "")[:120]])
            if rows:
                p = self.para("Supporting artifacts", bold=True, size=9, color="1F3864", space_after=2)
                U.keep_with_next(p)
                self.table(["Artifact", "Time (UTC)", "Summary", "Source"], rows, [3.2, 2.7, 6.6, 4.5], font_size=7.5)
        if f.get("mitre"):
            self.para("MITRE ATT&CK: " + ", ".join(f["mitre"]), italic=True, size=8.5, color="595959")

    def timeline_section(self):
        self.heading("6. Timeline of key events", 1)
        rows = self.db.query("SELECT * FROM timeline WHERE flagged=1 ORDER BY ts LIMIT 400")
        if not rows:
            rows = self.db.query("SELECT * FROM timeline ORDER BY ts LIMIT 400")
        if not rows:
            self.para("No timeline events were produced.")
            return
        self.para(f"The table lists the {len(rows)} events flagged by the analyzers in chronological order (UTC, with local time "
                  f"in {self.info.display_timezone}). The complete timeline is in the accompanying Excel workbook.")
        data = []
        for r in rows:
            loc = fmt(r["ts"], self.tz, with_zone=False) if self.tz != UTC else ""
            data.append([fmt(r["ts"], UTC, with_zone=False), loc, self.ev_label(r["evidence_id"]).split(" (")[0], r["source"],
                         r["event"], (r["description"] or "")[:180]])
        self.table(["UTC", "Local", "System", "Source", "Event", "Description"], data, [2.5, 2.5, 2.1, 1.8, 2.8, 5.3], font_size=7)

    def keywords_section(self):
        self.heading("7. Keyword and target-file search results", 1)
        for e in self.evidence:
            summ = self.db.artifacts(e["id"], "keyword_summary", order="id")
            matches = self.db.artifacts(e["id"], "target_match", order="ts")
            if not summ and not matches:
                continue
            self.heading(f"7.{self.evidence.index(e) + 1} {self.ev_label(e['id'])}", 2)
            if summ:
                rows = [[s["data"].get("term"), s["data"].get("label"), s["data"].get("total_hits"), (s["data"].get("areas") or "")[:220]]
                        for s in summ]
                self.table(["Search term", "Origin", "Hits", "Hits by area"], rows, [4.0, 2.6, 1.2, 9.2], font_size=7.5)
            if matches:
                rows = [[m["data"].get("target"), m["data"].get("match"), m["data"].get("location"), m["data"].get("area"),
                         m["data"].get("confidence")] for m in matches[:150]]
                self.table(["Target", "Match", "Location", "Area", "Confidence"], rows, [3.2, 3.4, 6.2, 2.6, 1.6], font_size=7.5)

    def coverage_section(self):
        self.heading("8. Artifact coverage - every location examined", 1)
        self.para("This matrix documents every artifact source the profile examined on each evidence item and the result, so that a "
                  "'no evidence found' conclusion can be read together with the locations that were actually checked.")
        for e in self.evidence:
            rows = self.db.query("SELECT module, artifact, location, status, count, detail FROM coverage WHERE evidence_id=? "
                                 "ORDER BY id", (e["id"],))
            if not rows:
                continue
            self.heading(f"8.{self.evidence.index(e) + 1} {self.ev_label(e['id'])}", 2)
            data = [[r["artifact"], (r["location"] or "")[:140], r["status"], r["count"] or "", (r["detail"] or "")[:160]] for r in rows]
            self.table(["Artifact", "Location checked", "Result", "Records", "Notes"], data, [3.8, 5.2, 2.6, 1.4, 4.0], font_size=7,
                       status_col=2, status_map=U.COV_COLORS)

    def limitations(self):
        self.heading("9. Limitations and assumptions", 1)
        items = []
        for r in self.db.query("SELECT evidence_id, artifact, location, detail FROM coverage WHERE status='error'"):
            items.append(f"{self.ev_label(r['evidence_id'])}: {r['artifact']} ({r['location']}) could not be processed - {r['detail']}.")
        for r in self.db.query("SELECT evidence_id, location, detail FROM coverage WHERE detail LIKE '%compressed%' OR detail LIKE '%BitLocker%'"):
            items.append(f"{self.ev_label(r['evidence_id'])}: {r['location']} - {r['detail']}.")
        for v in self.db.query("SELECT evidence_id, name, info_json FROM volumes WHERE fs LIKE 'FAT%'"):
            info = json.loads(v["info_json"] or "{}")
            if info.get("time_zone_assumed"):
                items.append(f"{self.ev_label(v['evidence_id'])} {v['name']}: FAT timestamps are stored in local time without a time "
                             f"zone; they were interpreted as {info['time_zone_assumed']}.")
        for e in self.evidence:
            if not (e.get("hashes") or {}).get("md5"):
                items.append(f"{self.ev_label(e['id'])}: the image hash was not re-verified in this run (container integrity checks only).")
        items += [
            "Registry key last-write times record the last change to a key, not necessarily the event of interest; MRU lists only "
            "time-stamp their most recent entry.",
            "NTFS last-access timestamps may be disabled or updated lazily by Windows; they are used as supporting indicators only.",
            "Browser history shows pages visited; it does not by itself prove that a file was uploaded. Upload conclusions rely on "
            "corroboration (file dialogs, memory artifacts, DLP alerts).",
            "Absence of an artifact does not prove that an action did not take place - artifacts can be overwritten, rotated or "
            "deliberately removed.",
        ]
        for i in items:
            self.bullet(i)

    def appendix(self):
        self.heading("Appendix A - Exported and carved files", 1)
        rows = []
        for r in self.db.query("SELECT evidence_id, source_path, local_path, size, sha256 FROM exported ORDER BY id LIMIT 300"):
            rows.append([self.ev_label(r["evidence_id"]).split(" (")[0], r["source_path"], os.path.basename(r["local_path"]), r["size"],
                         (r["sha256"] or "")[:64]])
        for a in self.db.artifacts(type_="carved_file", order="id", limit=300):
            d = a["data"]
            if d.get("saved_as"):
                rows.append([self.ev_label(a["evidence_id"]).split(" (")[0], f"carved {d.get('volume')} @ {d.get('offset'):#x}",
                             os.path.basename(d["saved_as"]), d.get("size"), d.get("sha256")])
        if rows:
            self.table(["Evidence", "Source", "Saved as", "Size", "SHA-256"], rows, [2.4, 5.0, 3.4, 1.4, 4.8], font_size=7, mono_cols=(4,))
        else:
            self.para("No files were exported.")
        self.parsed_appendix()
        self.heading("Appendix C - Software", 1)
        import dissect.target
        import docx as _docx
        import PySide6

        vers = [(__app_title__, __version__), ("Python", platform.python_version()), ("dissect.target", getattr(dissect.target, "__version__", "")),
                ("python-docx", getattr(_docx, "__version__", "")), ("Qt (PySide6)", PySide6.__version__),
                ("Host", f"{platform.system()} {platform.release()}")]
        try:
            import pyewf

            vers.append(("libewf", pyewf.get_version()))
        except Exception:
            pass
        self.kv_table(vers)
        self.heading("Appendix D - Glossary of artifacts", 1)
        for k, v in GLOSSARY:
            self.bullet(v, bold_prefix=f"{k}: ")

    def parsed_appendix(self):
        from ..core.exporter import evidence_folder

        self.heading("Appendix B - Parsed artifacts and raw collection for manual analysis", 1)
        self.para("Every artifact parsed by the tool is available as CSV for independent review (open in Excel / Timeline Explorer), "
                  "and the original artifact files were copied out of the image with their hashes. Folder paths are relative to the "
                  f"case folder {self.case.path}.")
        for e in self.evidence:
            pdir = evidence_folder(self.case, e, "Parsed")
            cdir = evidence_folder(self.case, e, "Collected")
            rows = []
            for root, _dirs, files in os.walk(pdir):
                for fn in sorted(files):
                    if not fn.lower().endswith(".csv"):
                        continue
                    full = os.path.join(root, fn)
                    try:
                        with open(full, encoding="utf-8-sig", errors="replace") as fh:
                            n = max(0, sum(1 for _ in fh) - 1)
                    except OSError:
                        n = ""
                    rows.append([os.path.relpath(full, self.case.path), n, f"{os.path.getsize(full) / 1024:,.0f} KB"])
            if not rows and not os.path.isdir(cdir):
                continue
            self.heading(f"{self.ev_label(e['id'])}", 2)
            if rows:
                self.table(["Parsed file", "Rows", "Size"], rows[:250], [12.6, 1.8, 2.6], font_size=7)
            man = os.path.join(cdir, "_manifest.csv")
            if os.path.exists(man):
                import csv as _csv

                with open(man, encoding="utf-8-sig") as fh:
                    mrows = list(_csv.DictReader(fh))
                groups = {}
                for m in mrows:
                    g = groups.setdefault(m.get("Group") or "-", [0, 0])
                    if m.get("CollectedAs"):
                        g[0] += 1
                        g[1] += int(m.get("Size") or 0)
                self.para(f"Raw collection: {sum(g[0] for g in groups.values()):,} files, "
                          f"{sum(g[1] for g in groups.values()) / 1e6:,.1f} MB in {os.path.relpath(cdir, self.case.path)} "
                          "(hash manifest: _manifest.csv).", size=9)
                self.table(["Group", "Files", "Size"], [[k, v[0], f"{v[1] / 1e6:,.1f} MB"] for k, v in sorted(groups.items())],
                           [8.0, 3.0, 3.0], font_size=8)

    def signoff(self):
        self.heading("Examiner declaration", 1)
        self.para("I confirm that the examination described in this report was carried out by me or under my supervision, that the "
                  "evidence was handled in a manner that preserved its integrity, and that the findings are a true reflection of the "
                  "artifacts examined.")
        self.kv_table([("Examiner", self.info.examiner or ""), ("Signature", ""), ("Date", "")], (4.0, 13.0), "FFFFFF")

    # ================================================================== main
    def write(self, logo_path):
        for i, f in enumerate(sorted(self.findings, key=lambda f: f["sort_key"]), start=1):
            self.fnum[f["id"]] = f"F-{i:02d}"
        self.setup()
        self.progress(0.42, "Report: cover and summary")
        self.cover(logo_path)
        self.document_control()
        self.executive_summary()
        self.background()
        self.evidence_section()
        self.methodology()
        self.new_page()
        self.findings_section()
        self.progress(0.86, "Report: timeline, coverage, appendices")
        self.new_page()
        self.timeline_section()
        self.keywords_section()
        self.coverage_section()
        self.limitations()
        self.new_page()
        self.appendix()
        self.signoff()
        self.doc.core_properties.title = self.profile.title
        self.doc.core_properties.author = self.info.examiner or __app_title__
        self.doc.core_properties.subject = self.info.case_name or ""
        self.doc.core_properties.comments = f"Generated by {__app_title__} {__version__}"
        self.doc.save(self.out_path)
        return self.out_path


GLOSSARY = [
    ("MFT ($MFT)", "NTFS Master File Table; one record per file, including records of deleted files until they are reused."),
    ("$STANDARD_INFORMATION / $FILE_NAME", "Two sets of NTFS timestamps; discrepancies between them can indicate timestamp manipulation."),
    ("USN journal ($UsnJrnl:$J)", "NTFS change journal recording file creation, modification, renaming and deletion."),
    ("File slack", "Space between the end of a file and the end of its last cluster; may hold remnants of earlier data."),
    ("Unallocated space", "Clusters not assigned to any file; deleted file content remains here until overwritten."),
    ("pagefile.sys / hiberfil.sys / swapfile.sys", "Memory written to disk by Windows; retains fragments of documents, URLs and commands."),
    ("Volume Shadow Copy", "Point-in-time snapshot of a volume kept by Windows; can contain earlier versions of deleted files."),
    ("LNK (shortcut) file", "Created when a user opens a file; records the target path, its timestamps, size, drive type, volume serial "
                            "number and volume label."),
    ("Jump list", "Per-application list of recently / frequently opened files with access times and counts."),
    ("Shellbags", "Registry records of folders viewed in Explorer, including folders on removable and network drives."),
    ("RecentDocs / OpenSavePidlMRU / LastVisitedPidlMRU", "Registry lists of recently opened documents, files chosen in Open/Save "
                                                          "dialogs and the folder each application last used in such a dialog."),
    ("RunMRU", "Commands typed into the Windows Run (Win+R) box - the primary artifact for ClickFix / paste-and-run attacks."),
    ("UserAssist", "Registry record of programs started through the Explorer GUI, with run counts and last run time."),
    ("Prefetch", "Files created when a program runs; record up to 8 run times and the files the program loaded."),
    ("Amcache / Shimcache / BAM", "Program inventory and compatibility caches that record executables present on or run by the system."),
    ("USBSTOR / MountedDevices / MountPoints2", "Registry records of USB storage devices, the drive letters / volumes they were mounted "
                                                "as, and the user profiles that mounted them."),
    ("Partition/Diagnostic event 1006", "Windows 10+ event logged at every disk arrival and removal, including the device serial and the "
                                        "first sector of each volume (volume serial number and label)."),
    ("setupapi.dev.log", "Driver installation log recording the first time each device was installed."),
    ("Zone.Identifier", "Mark-of-the-Web alternate data stream recording the URL a file was downloaded from."),
    ("Windows Timeline (ActivitiesCache.db)", "Database of applications and documents used, with focus start and end times."),
    ("SRUM", "System Resource Usage Monitor; per-application network bytes sent / received in hourly intervals."),
    ("Carving", "Recovery of files from unallocated space by recognizing file signatures and structures."),
]


# ====================================================================== public entry point
def render_figures(case, progress=None, cancel_event=None) -> dict:
    from .figures import ensure_app, render

    ensure_app()
    out = case.sub("figures")
    for f in os.listdir(out):
        if f.endswith(".png"):
            try:
                os.remove(os.path.join(out, f))
            except OSError:
                pass
    findings = case.db.findings(include_only=True)
    specs = [(f["id"], i, fig) for f in findings for i, fig in enumerate(f.get("figures") or [])]
    paths = {}
    for n, (fid, i, fig) in enumerate(specs):
        if cancel_event is not None and cancel_event.is_set():
            break
        p = os.path.join(out, f"finding{fid:03d}_{i}.png")
        try:
            render(fig, p)
            paths[(fid, i)] = p
        except Exception as e:  # pragma: no cover - logged, report continues
            import logging

            logging.getLogger("winforensics.report").exception("figure %s/%s failed: %s", fid, i, e)
        if progress:
            progress(0.05 + 0.3 * (n + 1) / max(1, len(specs)), f"Rendering figure {n + 1}/{len(specs)}")
    return paths


def finalize_with_word(docx_path: str, pdf: bool = True) -> str | None:
    """Use Microsoft Word (if installed) to update the TOC / page fields and export a PDF."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=False, AddToRecentFiles=False)
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Fields.Update()
        doc.Save()
        pdf_path = None
        if pdf:
            pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
            doc.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)
        doc.Close(False)
        return pdf_path or docx_path
    except Exception:
        return None
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def build_report(case, progress=None, cancel_event=None) -> dict:
    from .brand import logo_png
    from .workbook import build_workbook

    progress = progress or (lambda f, s: None)
    stamp = datetime.now(UTC).strftime("%Y%m%d_%H%M")
    base = re.sub(r"[^A-Za-z0-9_.-]+", "_", case.info.case_number or case.title)[:60]
    out_dir = case.sub("reports")
    progress(0.02, "Rendering figures")
    figs = render_figures(case, progress, cancel_event)
    progress(0.36, "Writing Excel workbook")
    xlsx = build_workbook(case, os.path.join(out_dir, f"{base}_artifacts_{stamp}.xlsx"), progress=lambda f, s: progress(0.36 + 0.06 * f, s))
    logo = logo_png(os.path.join(case.sub("figures"), "logo.png"))
    docx_path = os.path.join(out_dir, f"{base}_report_{stamp}.docx")
    ReportWriter(case, figs, docx_path, progress).write(logo)
    out = {"docx": docx_path, "xlsx": xlsx, "figures": case.sub("figures")}
    if (case.info.report or {}).get("word_finalize", True):
        progress(0.92, "Updating table of contents and exporting PDF (Microsoft Word)")
        res = finalize_with_word(docx_path, pdf=(case.info.report or {}).get("pdf", True))
        if res and res.endswith(".pdf"):
            out["pdf"] = res
        if not res:
            d = Document(docx_path)
            U.update_fields_on_open(d)
            d.save(docx_path)
    out["summary"] = write_summary_txt(case, os.path.join(out_dir, f"{base}_summary_{stamp}.txt"), out)
    progress(1.0, "Report complete")
    case.db.meta("last_report", out)
    return out


def write_summary_txt(case, path: str, outputs: dict | None = None) -> str:
    """Plain-text case summary for tickets / e-mail: questions and conclusions, findings, where the outputs are."""
    from ..profiles import get_profile

    db = case.db
    info = case.info
    try:
        prof = get_profile(info.profile)
        questions = prof.questions
        pname = prof.name
    except Exception:
        questions, pname = [], info.profile
    answers = db.answers()
    lines = [f"{info.case_number or ''} {info.case_name or case.title}".strip(), "=" * 78,
             f"Profile : {pname}", f"Examiner: {info.examiner or '-'}    Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M')} UTC", ""]
    lines.append("EVIDENCE")
    for e in db.evidence():
        osd = e.get("os") or {}
        lines.append(f"  E{e['id']:02d} {e['label']} - {osd.get('hostname') or '-'} {osd.get('product_name') or ''} ({e.get('format')})")
    lines += ["", "QUESTIONS AND CONCLUSIONS"]
    for i, q in enumerate(questions, 1):
        a = answers.get(q["id"], {})
        lines.append(f"  Q{i}. {q['text']}")
        lines.append(f"      -> {a.get('status', 'Inconclusive')}")
        for s in [x.strip() for x in (a.get("summary") or "").split("\n") if x.strip()][:8]:
            lines.append(f"         - {s[:300]}")
    lines += ["", "FINDINGS"]
    for f in db.findings():
        if f.get("include") == 0:
            continue
        lines.append(f"  - [{f.get('category') or '-'}] {f['title']}" + (f" ({str(f['ts'])[:19]} UTC)" if f.get("ts") else ""))
    lines += ["", "OUTPUTS"]
    for k, v in (outputs or {}).items():
        lines.append(f"  {k:8}: {v}")
    lines.append(f"  parsed  : {os.path.join(case.path, 'Parsed')}  (CSV of every artifact, SuperTimeline.csv)")
    lines.append(f"  raw     : {os.path.join(case.path, 'Collected')}  (original artifact files + hash manifest)")
    lines += ["", "All findings were generated from the examined evidence by automated analysis and must be reviewed by the examiner."]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return path
